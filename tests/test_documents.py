from io import BytesIO
import pymupdf as fitz
from docx import Document
from doctrad2.extraction import extract
from doctrad2.reconstruction import build, validate_pdf
from doctrad2.storage import Artifacts
from conftest import pdf


def test_native_graphics_links_bookmarks_preserved():
    source = pdf()
    ir = extract(source, "pdf")
    assert len(ir.nodes) == 1
    target, report = build(source, ir, {ir.nodes[0].id: "Tension d’entrée 12 V"})
    report = validate_pdf(source, target, report)
    assert not report.blocked, report
    assert report.metrics["outside_text_pixel_change_ratio"] == 0
    with fitz.open(stream=source, filetype="pdf") as a, fitz.open(stream=target, filetype="pdf") as b:
        assert len(a[0].get_drawings()) == len(b[0].get_drawings())
        assert a.get_toc() == b.get_toc()
        assert a[0].get_links()[0]["uri"] == b[0].get_links()[0]["uri"]
        assert "12 V" in b[0].get_text()
        assert "Input voltage" not in b[0].get_text()


def test_overflow_blocked_and_source_retained():
    source = pdf("12 V")
    ir = extract(source, "pdf")
    target, report = build(source, ir, {ir.nodes[0].id: "Une phrase qui ne tient pas dans la cellule " * 30})
    assert report.blocked
    assert any(f.code == "overflow" for f in report.findings)
    with fitz.open(stream=target, filetype="pdf") as doc:
        assert "12 V" in doc[0].get_text()


def test_ir_ids_stable_and_schema_roundtrip():
    from doctrad2.domain import DocumentIR
    source = pdf()
    a, b = extract(source, "pdf"), extract(source, "pdf")
    assert a.nodes[0].id == b.nodes[0].id
    assert DocumentIR.model_validate_json(a.model_dump_json()) == a


def test_rotated_page_uses_canonical_coordinates():
    doc = fitz.open(stream=pdf(), filetype="pdf")
    doc[0].set_rotation(90)
    ir = extract(doc.tobytes(), "pdf")
    assert not any(f.code == "rotated_text" for f in ir.findings)
    assert ir.pages[0].rotation == 90
    target, report = build(doc.tobytes(), ir, {ir.nodes[0].id: "Tension d’entrée 12 V"})
    assert not report.blocked
    with fitz.open(stream=target, filetype="pdf") as out:
        assert out[0].rotation == 90


def test_docx_preserves_mixed_runs_and_graphic_relationships():
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Input ").bold = True
    paragraph.add_run("voltage 12 V").italic = True
    output = BytesIO()
    doc.save(output)
    source = output.getvalue()
    ir = extract(source, "docx")
    translated, report = build(source, ir, {ir.nodes[0].id: "Tension d’entrée 12 V"})
    result = Document(BytesIO(translated))
    assert result.paragraphs[0].text == "Tension d’entrée 12 V"
    assert result.paragraphs[0].runs[0].bold
    assert result.paragraphs[0].runs[1].italic
    assert all(r.text for r in result.paragraphs[0].runs)


def test_artifacts_hash_integrity_and_path_safety(tmp_path):
    import pytest
    store = Artifacts(tmp_path)
    sha = store.put(b"same document")
    assert sha == store.put(b"same document")
    assert store.get(sha) == b"same document"
    with pytest.raises(ValueError):
        store.path("../secret")
    store.path(sha).write_bytes(b"corruption")
    with pytest.raises(ValueError):
        store.get(sha)


def test_orthogonal_text_rotation_preserved():
    doc = fitz.open()
    page = doc.new_page(width=600, height=800)
    page.insert_text((100, 400), "Input voltage 12 V", fontsize=12, rotate=90)
    source = doc.tobytes()
    ir = extract(source, "pdf")
    assert ir.nodes[0].locator["text_rotation"] == 90
    output, report = build(source, ir, {ir.nodes[0].id: "Tension 12 V"})
    assert not report.blocked
    with fitz.open(stream=output, filetype="pdf") as translated:
        assert "12 V" in translated[0].get_text()


def test_scan_without_ocr_is_blocked():
    original = fitz.open(stream=pdf(), filetype="pdf")
    image = original[0].get_pixmap().tobytes("png")
    scan = fitz.open()
    page = scan.new_page(width=600, height=800)
    page.insert_image(page.rect, stream=image)
    ir = extract(scan.tobytes(), "pdf")
    assert ir.pages[0].scanned
    assert any(f.code == "ocr_required" and f.severity == "critical" for f in ir.findings)


def test_large_false_table_does_not_swallow_entire_page():
    doc = fitz.open()
    page = doc.new_page(width=600, height=800)
    page.draw_rect(page.rect)
    page.draw_line((300, 0), (300, 800))
    page.draw_line((0, 400), (600, 400))
    page.insert_text((50, 100), "First paragraph on a decorated page", fontsize=12)
    page.insert_text((50, 160), "Second independent paragraph", fontsize=12)
    ir = extract(doc.tobytes(), "pdf")
    assert all(n.kind != "table_cell" for n in ir.nodes)
    assert len(ir.nodes) == 2


def test_link_intersecting_translated_text_survives_redaction():
    doc = fitz.open()
    page = doc.new_page(width=600, height=800)
    page.insert_text((60, 100), "Visit our website", fontsize=12)
    box = page.search_for("Visit our website")[0]
    page.insert_link({"kind": fitz.LINK_URI, "from": box, "uri": "https://example.org"})
    source = doc.tobytes()
    ir = extract(source, "pdf")
    result, report = build(source, ir, {ir.nodes[0].id: "Consultez notre site"})
    assert not any(f.code == "links" for f in report.findings)
    with fitz.open(stream=result, filetype="pdf") as reopened:
        assert reopened[0].get_links()[0]["uri"] == "https://example.org"
        assert "Visit our website" not in reopened[0].get_text()
